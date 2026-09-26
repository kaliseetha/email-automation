from __future__ import annotations

import json
import os
import re
import tempfile
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import streamlit as st
from openpyxl import Workbook, load_workbook


SHEET_HEADERS = {
    "Candidates": [
        "candidate_id",
        "name",
        "role",
        "email",
        "phone",
        "experience",
        "location",
        "skills",
        "profile_details",
        "created_at",
    ],
    "Accounts": [
        "account_id",
        "account_name",
        "to_emails",
        "cc_emails",
        "active",
        "created_at",
    ],
    "Submissions": [
        "tracking_id",
        "campaign_id",
        "candidate_id",
        "candidate_name",
        "account_id",
        "account_name",
        "subject",
        "to_emails",
        "cc_emails",
        "status",
        "sent_at",
        "response_at",
        "responder_name",
        "responder_email",
        "response_subject",
    ],
}

EMAIL_PATTERN = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+$")
REFERENCE_PATTERN = re.compile(r"\[REF:([0-9a-f]{12})\]", re.IGNORECASE)
DEMO_SENT_STATUSES = {"Sent", "Demo sent"}
APP_DIR = Path(__file__).resolve().parent
WORKBOOK_PATH = APP_DIR / "candidate_tracker.xlsx"
DEMO_MAILBOX_PATH = APP_DIR / "demo_outlook.json"


class WorkbookStore:
    def __init__(self, path: Path) -> None:
        self.path = path
        self._ensure_workbook()

    def _ensure_workbook(self) -> None:
        if self.path.exists():
            workbook = load_workbook(self.path)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)

        changed = False
        for sheet_name, expected_headers in SHEET_HEADERS.items():
            if sheet_name not in workbook.sheetnames:
                sheet = workbook.create_sheet(sheet_name)
                sheet.append(expected_headers)
                changed = True
                continue

            sheet = workbook[sheet_name]
            actual_headers = [cell.value for cell in sheet[1]]
            if actual_headers != expected_headers:
                raise ValueError(
                    f"The '{sheet_name}' sheet in {self.path.name} has unexpected "
                    "columns. Use a workbook created by this app or restore its "
                    "original column headers."
                )

        if changed or not self.path.exists():
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(self.path)
        workbook.close()

    def rows(self, sheet_name: str) -> list[dict[str, Any]]:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in sheet[1]]
            return [
                dict(zip(headers, values))
                for values in sheet.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in values)
            ]
        finally:
            workbook.close()

    def append(self, sheet_name: str, record: dict[str, Any]) -> None:
        workbook = load_workbook(self.path)
        try:
            sheet = workbook[sheet_name]
            headers = SHEET_HEADERS[sheet_name]
            sheet.append([record.get(header) for header in headers])
            workbook.save(self.path)
        finally:
            workbook.close()

    def update_submission(self, tracking_id: str, updates: dict[str, Any]) -> None:
        workbook = load_workbook(self.path)
        try:
            sheet = workbook["Submissions"]
            headers = [cell.value for cell in sheet[1]]
            tracking_column = headers.index("tracking_id") + 1
            update_columns = {
                headers.index(key) + 1: value for key, value in updates.items()
            }
            for row_number in range(2, sheet.max_row + 1):
                if sheet.cell(row=row_number, column=tracking_column).value == tracking_id:
                    for column, value in update_columns.items():
                        sheet.cell(row=row_number, column=column, value=value)
                    workbook.save(self.path)
                    return
            raise KeyError(f"Submission {tracking_id} was not found in the workbook.")
        finally:
            workbook.close()


class DemoMailboxStore:
    def __init__(self, path: Path) -> None:
        self.path = path
        if not self.path.exists():
            self._write({"sent_messages": [], "inbox_messages": []})

    def _read(self) -> dict[str, list[dict[str, Any]]]:
        try:
            with self.path.open("r", encoding="utf-8") as mailbox_file:
                data = json.load(mailbox_file)
        except json.JSONDecodeError as exc:
            raise ValueError(
                f"The demo mailbox file {self.path.name} contains invalid JSON. "
                "Restore a valid file or remove it to create a new empty demo mailbox."
            ) from exc

        if not isinstance(data, dict):
            raise ValueError(f"The demo mailbox file {self.path.name} has an invalid format.")
        for key in ("sent_messages", "inbox_messages"):
            if not isinstance(data.get(key), list):
                raise ValueError(
                    f"The demo mailbox file {self.path.name} is missing the '{key}' list."
                )
        return data

    def _write(self, data: dict[str, list[dict[str, Any]]]) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path: str | None = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                newline="\n",
                dir=self.path.parent,
                prefix=f".{self.path.name}.",
                suffix=".tmp",
                delete=False,
            ) as temporary_file:
                temporary_path = temporary_file.name
                json.dump(data, temporary_file, indent=2, ensure_ascii=False)
                temporary_file.write("\n")
            os.replace(temporary_path, self.path)
        except OSError:
            if temporary_path and os.path.exists(temporary_path):
                os.unlink(temporary_path)
            raise

    def record_sent_message(self, message: dict[str, Any]) -> None:
        mailbox = self._read()
        mailbox["sent_messages"].append(message)
        self._write(mailbox)

    def record_reply(self, reply: dict[str, Any]) -> None:
        mailbox = self._read()
        mailbox["inbox_messages"].append(reply)
        self._write(mailbox)

    def scan_inbox(self) -> list[dict[str, Any]]:
        mailbox = self._read()
        matches = []
        for message in mailbox["inbox_messages"]:
            subject = str(message.get("subject", "") or "")
            references = REFERENCE_PATTERN.findall(subject)
            if not references:
                continue
            try:
                received_at = datetime.fromisoformat(str(message["received_at"]))
            except (KeyError, TypeError, ValueError) as exc:
                raise ValueError(
                    "A demo inbox message has an invalid received_at timestamp."
                ) from exc
            matches.append(
                {
                    "tracking_id": references[0].lower(),
                    "response_at": received_at,
                    "responder_name": str(message.get("sender_name", "") or ""),
                    "responder_email": str(message.get("sender_email", "") or ""),
                    "response_subject": subject[:500],
                }
            )
        return sorted(matches, key=lambda match: match["response_at"])


def parse_addresses(value: str | None) -> list[str]:
    return [
        address.strip()
        for address in re.split(r"[,;]", value or "")
        if address.strip()
    ]


def tracker_error_message(exc: OSError, action: str) -> str:
    if isinstance(exc, PermissionError):
        return (
            f"Could not {action}: the Excel tracker is locked or not writable. "
            f"Close '{WORKBOOK_PATH.name}' in Excel (and any other program using it), "
            "then retry. The change was not saved. Make sure the app folder is writable."
        )
    return f"Could not {action}: {exc}"


def validate_addresses(
    value: str | None, field_name: str, required: bool = False
) -> list[str]:
    addresses = parse_addresses(value)
    if required and not addresses:
        raise ValueError(f"Enter at least one email address in {field_name}.")
    invalid = [address for address in addresses if not EMAIL_PATTERN.fullmatch(address)]
    if invalid:
        raise ValueError(f"Invalid address(es) in {field_name}: {', '.join(invalid)}")
    return addresses


def candidate_message(candidate: dict[str, Any]) -> str:
    fields = [
        ("Candidate", candidate.get("name")),
        ("Role", candidate.get("role")),
        ("Email", candidate.get("email")),
        ("Phone", candidate.get("phone")),
        ("Experience", candidate.get("experience")),
        ("Location", candidate.get("location")),
        ("Skills", candidate.get("skills")),
        ("Profile details", candidate.get("profile_details")),
    ]
    return "\n".join(
        f"{label}: {value}" for label, value in fields if value not in (None, "")
    )


def format_datetime(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value or "")


def first_responses(submissions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    campaigns: dict[str, list[dict[str, Any]]] = {}
    for submission in submissions:
        if submission.get("response_at"):
            campaigns.setdefault(str(submission.get("campaign_id")), []).append(submission)

    results = []
    for responses in campaigns.values():
        first = min(responses, key=lambda row: row["response_at"])
        results.append(
            {
                "Candidate": first.get("candidate_name"),
                "First responding account": first.get("account_name"),
                "Responder": first.get("responder_name"),
                "Responder email": first.get("responder_email"),
                "Received at": format_datetime(first.get("response_at")),
                "Campaign": first.get("campaign_id"),
            }
        )
    return sorted(results, key=lambda row: row["Received at"], reverse=True)


def render_candidate_form(store: WorkbookStore) -> None:
    st.subheader("Add candidate")
    with st.form("candidate_form", clear_on_submit=True):
        name = st.text_input("Candidate name *")
        role = st.text_input("Role / position *")
        email = st.text_input("Candidate email")
        phone = st.text_input("Phone")
        experience = st.text_input("Experience (for example, 5 years)")
        location = st.text_input("Location")
        skills = st.text_area("Skills")
        profile_details = st.text_area("Additional profile details")
        submitted = st.form_submit_button("Save candidate")

    if submitted:
        if not name.strip() or not role.strip():
            st.error("Candidate name and role are required.")
        elif email.strip() and not EMAIL_PATTERN.fullmatch(email.strip()):
            st.error("Enter a valid candidate email address.")
        else:
            try:
                store.append(
                    "Candidates",
                    {
                        "candidate_id": uuid.uuid4().hex[:12],
                        "name": name.strip(),
                        "role": role.strip(),
                        "email": email.strip(),
                        "phone": phone.strip(),
                        "experience": experience.strip(),
                        "location": location.strip(),
                        "skills": skills.strip(),
                        "profile_details": profile_details.strip(),
                        "created_at": datetime.now(),
                    },
                )
                st.success("Candidate saved to the Excel workbook.")
            except OSError as exc:
                st.error(tracker_error_message(exc, "save the candidate"))


def render_candidate_import(store: WorkbookStore) -> None:
    st.subheader("Import candidates from an existing Excel file")
    uploaded_file = st.file_uploader(
        "Choose an .xlsx or .xlsm file",
        type=["xlsx", "xlsm"],
        key="candidate_import_file",
    )
    if uploaded_file is None:
        return

    try:
        source_workbook = load_workbook(
            BytesIO(uploaded_file.getvalue()), read_only=True, data_only=True
        )
        sheet_name = st.selectbox(
            "Worksheet", options=source_workbook.sheetnames, key="candidate_import_sheet"
        )
        source_sheet = source_workbook[sheet_name]
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in source_sheet[1]
        ]
        if not headers or not any(headers):
            source_workbook.close()
            st.error("The selected worksheet has no column headers in its first row.")
            return

        aliases = {
            "name": {"name", "candidate", "candidate name", "full name"},
            "role": {"role", "position", "job title", "designation"},
            "email": {"email", "email address", "candidate email"},
            "phone": {"phone", "mobile", "contact", "phone number"},
            "experience": {"experience", "years of experience", "total experience"},
            "location": {"location", "city", "current location"},
            "skills": {"skills", "technical skills", "key skills"},
            "profile_details": {"profile details", "details", "summary", "comments"},
        }
        st.caption(
            "Map the columns from your existing workbook. Unmapped columns are retained "
            "in Additional profile details."
        )
        column_options = ["(Skip)"] + headers
        mapping: dict[str, str] = {}
        with st.form("candidate_import_form"):
            for field, label in [
                ("name", "Candidate name"),
                ("role", "Role / position"),
                ("email", "Candidate email"),
                ("phone", "Phone"),
                ("experience", "Experience"),
                ("location", "Location"),
                ("skills", "Skills"),
                ("profile_details", "Additional profile details"),
            ]:
                default_header = next(
                    (
                        index + 1
                        for index, header in enumerate(headers)
                        if header.lower() in aliases[field]
                    ),
                    0,
                )
                selected = st.selectbox(
                    f"{label} column",
                    options=column_options,
                    index=default_header,
                    key=f"candidate_import_{field}",
                )
                mapping[field] = "" if selected == "(Skip)" else selected
            import_submitted = st.form_submit_button("Import candidate rows")

        if import_submitted:
            if not mapping["name"]:
                source_workbook.close()
                st.error("Map a column to Candidate name before importing.")
                return

            existing = store.rows("Candidates")
            existing_keys = {
                (
                    str(row.get("name") or "").strip().casefold(),
                    str(row.get("email") or "").strip().casefold(),
                )
                for row in existing
                if row.get("email")
            }
            imported = 0
            skipped = 0
            for values in source_sheet.iter_rows(min_row=2, values_only=True):
                source_record = {
                    header: values[index] if index < len(values) else None
                    for index, header in enumerate(headers)
                    if header
                }
                name = str(source_record.get(mapping["name"], "") or "").strip()
                if not name:
                    skipped += 1
                    continue
                email = str(source_record.get(mapping["email"], "") or "").strip()
                key = (name.casefold(), email.casefold())
                if email and key in existing_keys:
                    skipped += 1
                    continue

                profile_details = str(
                    source_record.get(mapping["profile_details"], "") or ""
                ).strip()
                mapped_headers = {header for header in mapping.values() if header}
                extras = [
                    f"{header}: {value}"
                    for header, value in source_record.items()
                    if header not in mapped_headers and value not in (None, "")
                ]
                if extras:
                    profile_details = "\n".join(
                        part for part in (profile_details, "\n".join(extras)) if part
                    )

                imported_candidate = {
                    "candidate_id": uuid.uuid4().hex[:12],
                    "name": name,
                    "role": str(
                        source_record.get(mapping["role"], "Not specified")
                        or "Not specified"
                    ).strip(),
                    "email": email,
                    "phone": str(source_record.get(mapping["phone"], "") or "").strip(),
                    "experience": str(
                        source_record.get(mapping["experience"], "") or ""
                    ).strip(),
                    "location": str(
                        source_record.get(mapping["location"], "") or ""
                    ).strip(),
                    "skills": str(source_record.get(mapping["skills"], "") or "").strip(),
                    "profile_details": profile_details,
                    "created_at": datetime.now(),
                }
                store.append("Candidates", imported_candidate)
                if email:
                    existing_keys.add(key)
                imported += 1
            source_workbook.close()
            st.success(f"Imported {imported} candidate(s); skipped {skipped} blank or duplicate row(s).")
    except (OSError, ValueError, KeyError) as exc:
        if isinstance(exc, OSError):
            st.error(tracker_error_message(exc, "import candidates"))
        else:
            st.error(f"Could not read or import this Excel file: {exc}")


def render_account_form(store: WorkbookStore) -> None:
    st.subheader("Add account")
    with st.form("account_form", clear_on_submit=True):
        account_name = st.text_input("Account / client name *")
        to_emails = st.text_input("To email(s) *", help="Separate addresses with commas or semicolons.")
        cc_emails = st.text_input("CC email(s)", help="Optional; separate addresses with commas or semicolons.")
        submitted = st.form_submit_button("Save account")

    if submitted:
        try:
            validate_addresses(to_emails, "To", required=True)
            validate_addresses(cc_emails, "CC")
            if not account_name.strip():
                raise ValueError("Account / client name is required.")
            store.append(
                "Accounts",
                {
                    "account_id": uuid.uuid4().hex[:12],
                    "account_name": account_name.strip(),
                    "to_emails": "; ".join(parse_addresses(to_emails)),
                    "cc_emails": "; ".join(parse_addresses(cc_emails)),
                    "active": True,
                    "created_at": datetime.now(),
                },
            )
            st.success("Account saved to the Excel workbook.")
        except ValueError as exc:
            st.error(str(exc))
        except OSError as exc:
            st.error(tracker_error_message(exc, "save the account"))


def render_send_tab(store: WorkbookStore, mailbox: DemoMailboxStore) -> None:
    st.subheader("Send a candidate profile")
    st.info(
        "Demo mode: emails are recorded in demo_outlook.json. Nothing is sent through Outlook."
    )
    candidates = store.rows("Candidates")
    accounts = [account for account in store.rows("Accounts") if account.get("active")]
    if not candidates or not accounts:
        st.info("Add at least one candidate and one account before sending.")
        return

    candidate_by_id = {str(row["candidate_id"]): row for row in candidates}
    account_by_id = {str(row["account_id"]): row for row in accounts}
    selected_candidate_id = st.selectbox(
        "Candidate",
        options=list(candidate_by_id),
        format_func=lambda candidate_id: (
            f"{candidate_by_id[candidate_id].get('name')} - "
            f"{candidate_by_id[candidate_id].get('role')}"
        ),
    )
    selected_account_ids = st.multiselect(
        "Send individually to these accounts",
        options=list(account_by_id),
        format_func=lambda account_id: str(account_by_id[account_id].get("account_name")),
    )
    candidate = candidate_by_id[selected_candidate_id]
    subject = st.text_input(
        "Email subject",
        value=f"Candidate profile - {candidate.get('name')} - {candidate.get('role')}",
        key=f"email_subject_{selected_candidate_id}",
    )
    body = st.text_area(
        "Email body",
        value=candidate_message(candidate),
        height=260,
        key=f"email_body_{selected_candidate_id}",
    )
    st.caption(
        "A separate demo message is recorded for each account. Each subject gets a "
        "unique [REF:...] tag that the demo inbox uses to match replies."
    )

    if st.button("Record demo emails for selected accounts", type="primary"):
        if not selected_account_ids:
            st.error("Select at least one account.")
            return
        if not subject.strip() or not body.strip():
            st.error("Email subject and body are required.")
            return

        campaign_id = uuid.uuid4().hex[:12]
        results = []
        for account_id in selected_account_ids:
            account = account_by_id[account_id]
            tracking_id = uuid.uuid4().hex[:12]
            try:
                to_list = validate_addresses(account.get("to_emails", ""), "To", required=True)
                cc_list = validate_addresses(account.get("cc_emails", ""), "CC")
                tagged_subject = f"{subject.strip()} [REF:{tracking_id}]"
                store.append(
                    "Submissions",
                    {
                        "tracking_id": tracking_id,
                        "campaign_id": campaign_id,
                        "candidate_id": candidate.get("candidate_id"),
                        "candidate_name": candidate.get("name"),
                        "account_id": account.get("account_id"),
                        "account_name": account.get("account_name"),
                        "subject": tagged_subject,
                        "to_emails": "; ".join(to_list),
                        "cc_emails": "; ".join(cc_list),
                        "status": "Sending",
                    },
                )
            except Exception as exc:
                if isinstance(exc, OSError):
                    message = tracker_error_message(exc, "record the email")
                else:
                    message = str(exc)
                st.error(f"Could not prepare the email for {account.get('account_name')}: {message}")
                results.append((str(account.get("account_name")), False, str(exc)))
                continue

            try:
                sent_at = datetime.now()
                mailbox.record_sent_message(
                    {
                        "tracking_id": tracking_id,
                        "campaign_id": campaign_id,
                        "candidate_id": candidate.get("candidate_id"),
                        "candidate_name": candidate.get("name"),
                        "account_id": account.get("account_id"),
                        "account_name": account.get("account_name"),
                        "to_emails": to_list,
                        "cc_emails": cc_list,
                        "subject": tagged_subject,
                        "body": body,
                        "sent_at": sent_at.isoformat(timespec="seconds"),
                    }
                )
            except Exception as exc:
                try:
                    store.update_submission(tracking_id, {"status": "Send failed"})
                except Exception as status_exc:
                    st.error(
                        f"Recording the demo email failed for {account.get('account_name')}, "
                        f"and the tracker could not record that failure: "
                        f"{tracker_error_message(status_exc, 'update the tracker') if isinstance(status_exc, OSError) else status_exc}"
                    )
                message = (
                    f"Could not save the demo email to {DEMO_MAILBOX_PATH.name}: {exc}"
                )
                st.error(f"Could not record the demo email for {account.get('account_name')}: {message}")
                results.append((str(account.get("account_name")), False, str(exc)))
                continue

            try:
                store.update_submission(
                    tracking_id, {"status": "Demo sent", "sent_at": datetime.now()}
                )
            except Exception as exc:
                st.error(
                    f"The demo email was recorded for {account.get('account_name')}, "
                    f"but the Excel tracker could not be updated. Reference "
                    f"{tracking_id}: "
                    f"{tracker_error_message(exc, 'update the tracker') if isinstance(exc, OSError) else exc}"
                )
                results.append(
                    (
                        str(account.get("account_name")),
                        True,
                        "Demo email recorded; tracker update failed",
                    )
                )
                continue
            results.append((str(account.get("account_name")), True, "Demo email recorded"))

        for account_name, sent, message in results:
            if sent:
                st.success(f"{account_name}: {message}.")
        if any(result[1] for result in results):
            st.info(
                "Use the response-tracking tab to simulate replies and scan the demo inbox."
            )


def render_demo_reply_form(
    store: WorkbookStore, mailbox: DemoMailboxStore
) -> None:
    submissions = [
        row for row in store.rows("Submissions") if row.get("status") in DEMO_SENT_STATUSES
    ]
    st.markdown("#### Simulate an account reply")
    if not submissions:
        st.info("Send a demo candidate profile before simulating a reply.")
        return

    submission_by_id = {str(row["tracking_id"]): row for row in submissions}
    with st.form("demo_reply_form", clear_on_submit=True):
        tracking_id = st.selectbox(
            "Reply to a sent profile",
            options=list(submission_by_id),
            format_func=lambda value: (
                f"{submission_by_id[value].get('candidate_name')} — "
                f"{submission_by_id[value].get('account_name')} "
                f"({submission_by_id[value].get('subject')})"
            ),
        )
        responder_name = st.text_input("Responder name *")
        responder_email = st.text_input("Responder email *")
        reply_body = st.text_area("Reply message", value="We are interested in this candidate.")
        received_date = st.date_input("Received date", value=datetime.now().date())
        received_time = st.time_input(
            "Received time", value=datetime.now().time().replace(microsecond=0)
        )
        submitted = st.form_submit_button("Add demo reply to inbox")

    if not submitted:
        return
    if not responder_name.strip():
        st.error("Enter the responder's name.")
        return
    try:
        validate_addresses(responder_email, "Responder email", required=True)
        submission = submission_by_id[tracking_id]
        mailbox.record_reply(
            {
                "tracking_id": tracking_id,
                "sender_name": responder_name.strip(),
                "sender_email": responder_email.strip(),
                "subject": f"Re: {submission.get('subject')}",
                "body": reply_body.strip(),
                "received_at": datetime.combine(
                    received_date, received_time
                ).isoformat(timespec="seconds"),
            }
        )
    except ValueError as exc:
        st.error(str(exc))
    except OSError as exc:
        st.error(f"Could not save the demo reply to {DEMO_MAILBOX_PATH.name}: {exc}")
    else:
        st.success(
            f"Demo reply saved for {submission.get('account_name')}. Scan the demo inbox "
            "to update the tracker."
        )


def render_tracking_tab(store: WorkbookStore, mailbox: DemoMailboxStore) -> None:
    st.subheader("Find the first account to respond")
    st.write(
        "Demo mode reads replies from demo_outlook.json. Add simulated account replies "
        "below, then scan the file to identify the first responder for each campaign."
    )
    if st.button("Scan demo inbox file", type="primary"):
        try:
            inbox_matches = mailbox.scan_inbox()
            submissions = store.rows("Submissions")
            by_tracking_id = {
                str(row.get("tracking_id", "")).lower(): row for row in submissions
            }
            updated = 0
            for match in inbox_matches:
                submission = by_tracking_id.get(match["tracking_id"])
                if not submission or submission.get("status") not in DEMO_SENT_STATUSES:
                    continue
                existing_at = submission.get("response_at")
                if existing_at and existing_at <= match["response_at"]:
                    continue
                store.update_submission(
                    match["tracking_id"],
                    {
                        "response_at": match["response_at"],
                        "responder_name": match["responder_name"],
                        "responder_email": match["responder_email"],
                        "response_subject": match["response_subject"],
                    },
                )
                submission.update(match)
                updated += 1
            st.success(
                f"Demo inbox scan complete. Recorded or refreshed {updated} account response(s)."
            )
        except Exception as exc:
            st.error(f"Could not read demo inbox {DEMO_MAILBOX_PATH.name}: {exc}")

    submissions = store.rows("Submissions")
    if not submissions:
        st.info("No candidate profiles have been sent yet.")
        return

    st.markdown("#### First response per campaign")
    earliest = first_responses(submissions)
    if earliest:
        st.dataframe(earliest, width="stretch", hide_index=True)
    else:
        st.info("No matched replies have been found yet.")

    st.markdown("#### Account-by-account status")
    statuses = [
        {
            "Candidate": row.get("candidate_name"),
            "Account": row.get("account_name"),
            "Status": "Responded" if row.get("response_at") else row.get("status"),
            "Sent at": format_datetime(row.get("sent_at")),
            "Response received at": format_datetime(row.get("response_at")),
            "Responder": row.get("responder_name"),
            "Responder email": row.get("responder_email"),
            "Subject / reference": row.get("subject"),
        }
        for row in sorted(
            submissions,
            key=lambda item: format_datetime(item.get("sent_at")),
            reverse=True,
        )
    ]
    st.dataframe(statuses, width="stretch", hide_index=True)
    render_demo_reply_form(store, mailbox)


def main() -> None:
    st.set_page_config(page_title="Candidate Account Tracker", layout="wide")
    st.title("Candidate Account Email Tracker")
    st.markdown(
        "A simple workspace for resource management teams to organize candidate "
        "profiles, share each profile with multiple client accounts, and quickly see "
        "which account responds first. Prepare candidate emails, track account-by-account "
        "interest, and keep the hiring process moving from one place."
    )
    st.caption(
        "Demo mode: emails and replies are simulated and saved locally. No real emails "
        "are sent, and no Outlook mailbox is accessed."
    )

    try:
        store = WorkbookStore(WORKBOOK_PATH)
    except (OSError, PermissionError, ValueError) as exc:
        if isinstance(exc, OSError):
            st.error(tracker_error_message(exc, "open the Excel tracker"))
        else:
            st.error(
                f"Could not open the Excel tracker {WORKBOOK_PATH.name}: {exc}"
            )
        st.stop()

    try:
        mailbox = DemoMailboxStore(DEMO_MAILBOX_PATH)
    except (OSError, ValueError) as exc:
        st.error(f"Could not open the demo mailbox file {DEMO_MAILBOX_PATH.name}: {exc}")
        st.stop()

    send_tab, tracking_tab, candidates_tab, accounts_tab = st.tabs(
        ["Send candidate", "Track responses", "Candidates", "Accounts"]
    )
    with send_tab:
        render_send_tab(store, mailbox)
    with tracking_tab:
        render_tracking_tab(store, mailbox)
    with candidates_tab:
        render_candidate_form(store)
        render_candidate_import(store)
        st.markdown("#### Saved candidates")
        st.dataframe(store.rows("Candidates"), width="stretch", hide_index=True)
    with accounts_tab:
        render_account_form(store)
        st.markdown("#### Saved accounts")
        st.dataframe(store.rows("Accounts"), width="stretch", hide_index=True)

    st.divider()
    try:
        tracker_bytes = WORKBOOK_PATH.read_bytes()
    except OSError as exc:
        st.error(tracker_error_message(exc, "read the Excel tracker for download"))
    else:
        st.download_button(
            "Download Excel tracker",
            data=tracker_bytes,
            file_name=WORKBOOK_PATH.name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    try:
        mailbox_bytes = DEMO_MAILBOX_PATH.read_bytes()
    except OSError as exc:
        st.error(f"Could not read the demo mailbox file for download: {exc}")
    else:
        st.download_button(
            "Download demo mailbox JSON",
            data=mailbox_bytes,
            file_name=DEMO_MAILBOX_PATH.name,
            mime="application/json",
        )


if __name__ == "__main__":
    main()